"""Geração determinística dos documentos de referência (.docx/.pptx).

O gerador lê `data/knowledge/*.md` (só `status: approved`) e materializa um
par docx/pptx por documento, com metadados visíveis no corpo e no nome do
arquivo. Mesmo insumo → bytes idênticos; o round-trip (abrir com python-docx/
python-pptx e reconstruir as seções) é a garantia de que a ingestão via MinIO
enxerga os mesmos trechos do markdown.
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from utils.generate_reference_documents import (
    FIXED_DATETIME,
    ReferenceDocument,
    generate_documents,
    load_knowledge_documents,
    parse_docx_document,
    parse_pptx_document,
    resolve_bucket,
    upload_documents,
)

_APPROVED = """---
document_id: doc-rentabilidade
title: Como o banco calcula rentabilidade
version: v1
status: approved
locale: pt-BR
approved_by: area-negocios-financas (ficticio)
---

## O que é rentabilidade e como ela é definida

Rentabilidade é a taxa anual que mede o retorno da carteira.

## Fórmula de rentabilidade e anualização

A margem oficial já é anualizada em reais por ano.
"""

_DRAFT = """---
document_id: doc-rascunho
title: Rascunho de método
version: v1
status: draft
locale: pt-BR
---

## Seção que não deve sair

Conteúdo não aprovado.
"""


def _corpus(tmp_path: Path) -> Path:

    knowledge = tmp_path / "knowledge"
    knowledge.mkdir()
    (knowledge / "01-rentabilidade.md").write_text(_APPROVED, encoding="utf-8")
    (knowledge / "02-rascunho.md").write_text(_DRAFT, encoding="utf-8")

    return knowledge


class _FakeMinio:
    def __init__(self, existing: tuple[str, ...] = ()):

        self.existing = set(existing)
        self.created: list[str] = []
        self.puts: list[tuple[str, str, str]] = []
        self.metadata: list[dict] = []

    def bucket_exists(self, bucket: str) -> bool:

        return bucket in self.existing

    def make_bucket(self, bucket: str) -> None:

        self.created.append(bucket)
        self.existing.add(bucket)

    def fput_object(
        self,
        bucket: str,
        name: str,
        path: str,
        content_type: str | None = None,
        metadata: dict | None = None,
    ) -> None:

        self.puts.append((bucket, name, path))
        self.metadata.append(metadata or {})


def test_gera_um_par_por_documento_aprovado(tmp_path: Path) -> None:

    written = generate_documents(_corpus(tmp_path), tmp_path / "reference")

    assert sorted(path.name for path in written) == [
        "doc-rentabilidade__v1.docx",
        "doc-rentabilidade__v1.pptx",
    ]
    assert not (tmp_path / "reference" / "doc-rascunho__v1.docx").exists()


def test_bytes_sao_deterministicos(tmp_path: Path) -> None:

    knowledge = _corpus(tmp_path)
    first = generate_documents(knowledge, tmp_path / "reference-a")
    second = generate_documents(knowledge, tmp_path / "reference-b")

    assert [path.name for path in first] == [path.name for path in second]

    for path_a, path_b in zip(first, second, strict=False):
        assert path_a.read_bytes() == path_b.read_bytes()


def test_docx_abre_com_metadados_e_round_trip(tmp_path: Path) -> None:

    knowledge = _corpus(tmp_path)
    generate_documents(knowledge, tmp_path / "reference")
    data = (tmp_path / "reference" / "doc-rentabilidade__v1.docx").read_bytes()

    approved = load_knowledge_documents(knowledge)[0]

    assert parse_docx_document(data) == approved

    doc = Document(io.BytesIO(data))
    paragraphs = [paragraph.text for paragraph in doc.paragraphs]

    assert "Título: Como o banco calcula rentabilidade" in paragraphs
    assert "Versão: v1" in paragraphs
    assert "Documento: doc-rentabilidade" in paragraphs
    assert "Status: approved" in paragraphs
    assert doc.core_properties.created == FIXED_DATETIME
    assert doc.core_properties.modified == FIXED_DATETIME


def test_pptx_abre_com_metadados_e_round_trip(tmp_path: Path) -> None:

    knowledge = _corpus(tmp_path)
    generate_documents(knowledge, tmp_path / "reference")
    data = (tmp_path / "reference" / "doc-rentabilidade__v1.pptx").read_bytes()

    approved = load_knowledge_documents(knowledge)[0]

    assert parse_pptx_document(data) == approved

    presentation = Presentation(io.BytesIO(data))
    slides = list(presentation.slides)

    assert len(slides) == 1 + len(approved.sections)
    cover_text = "\n".join(
        shape.text_frame.text for shape in slides[0].shapes if shape.has_text_frame
    )

    assert "Título: Como o banco calcula rentabilidade" in cover_text
    assert "Documento: doc-rentabilidade" in cover_text
    assert "Status: approved" in cover_text
    # python-pptx devolve a data naive; o valor gravado é o mesmo fixo do docx.
    assert presentation.core_properties.created == FIXED_DATETIME.replace(tzinfo=None)


def test_upload_cria_bucket_e_envia_os_arquivos(tmp_path: Path) -> None:

    written = generate_documents(_corpus(tmp_path), tmp_path / "reference")
    client = _FakeMinio()

    uploaded = upload_documents(written, bucket="cantabria-docs", client=client)

    assert uploaded == 2
    assert client.created == ["cantabria-docs"]
    assert [name for _, name, _ in client.puts] == [
        "documents/doc-rentabilidade/v1/doc-rentabilidade.docx",
        "documents/doc-rentabilidade/v1/doc-rentabilidade.pptx",
    ]
    assert client.metadata[0] == {
        "document_id": "doc-rentabilidade",
        "version": "v1",
        "status": "approved",
    }


def test_upload_usa_bucket_existente(tmp_path: Path) -> None:

    written = generate_documents(_corpus(tmp_path), tmp_path / "reference")
    client = _FakeMinio(existing=("cantabria-docs",))

    upload_documents(written, bucket="cantabria-docs", client=client)

    assert client.created == []


def test_bucket_default_e_env(monkeypatch: pytest.MonkeyPatch) -> None:

    monkeypatch.delenv("MINIO_BUCKET", raising=False)
    assert resolve_bucket() == "cantabria-docs"

    monkeypatch.setenv("MINIO_BUCKET", "docs-teste")
    assert resolve_bucket() == "docs-teste"


def test_acervo_ausente_falha_alto(tmp_path: Path) -> None:

    with pytest.raises(FileNotFoundError):
        generate_documents(tmp_path / "nao-existe", tmp_path / "reference")


def test_documento_sem_id_aprovado_e_ignorado(tmp_path: Path) -> None:

    knowledge = _corpus(tmp_path)
    (knowledge / "03-sem-id.md").write_text(
        "---\ntitle: Sem id\nversion: v1\nstatus: approved\n---\n\n## Seção\n\nCorpo.",
        encoding="utf-8",
    )

    documents = load_knowledge_documents(knowledge)

    assert [document.document_id for document in documents] == ["doc-rentabilidade"]
    assert isinstance(documents[0], ReferenceDocument)
